import io
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from .db import load_accounts, load_transactions, load_pending_payments, load_budgets_vs_actual

def generate_excel_report():
    """Generates an Excel workbook (.xlsx) in memory with formatted sheets."""
    output = io.BytesIO()
    
    accounts_df = load_accounts()
    transactions_df = load_transactions(limit=1000)
    pending_df = load_pending_payments()
    budgets_df = load_budgets_vs_actual()

    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        # Sheet 1: Accounts
        if not accounts_df.empty:
            accounts_df.to_excel(writer, sheet_name='Cuentas y Saldos', index=False)
        else:
            pd.DataFrame({'Mensaje': ['Sin cuentas registradas']}).to_excel(writer, sheet_name='Cuentas y Saldos', index=False)

        # Sheet 2: Transactions
        if not transactions_df.empty:
            transactions_df.to_excel(writer, sheet_name='Transacciones', index=False)
        else:
            pd.DataFrame({'Mensaje': ['Sin transacciones registradas']}).to_excel(writer, sheet_name='Transacciones', index=False)

        # Sheet 3: Pending Payments
        if not pending_df.empty:
            pending_df.to_excel(writer, sheet_name='Pagos y Cobros Pendientes', index=False)
        else:
            pd.DataFrame({'Mensaje': ['Sin pagos pendientes']}).to_excel(writer, sheet_name='Pagos y Cobros Pendientes', index=False)

        # Sheet 4: Budgets
        if not budgets_df.empty:
            budgets_df.to_excel(writer, sheet_name='Presupuestos vs Real', index=False)
        else:
            pd.DataFrame({'Mensaje': ['Sin presupuestos registrados']}).to_excel(writer, sheet_name='Presupuestos vs Real', index=False)

    output.seek(0)
    
    # Apply openpyxl styling (Headers, column auto-fit)
    wb = openpyxl.load_workbook(output)
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )

    for sheetname in wb.sheetnames:
        ws = wb[sheetname]
        ws.views.sheetView[0].showGridLines = True
        
        # Style header row
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        # Auto-fit column widths
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                cell.border = thin_border
                val_str = str(cell.value or '')
                if len(val_str) > max_len:
                    max_len = len(val_str)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    final_output = io.BytesIO()
    wb.save(final_output)
    final_output.seek(0)
    return final_output.getvalue()
